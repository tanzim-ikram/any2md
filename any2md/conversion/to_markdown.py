"""Document → Markdown conversion using Microsoft MarkItDown."""

from __future__ import annotations

import traceback
from pathlib import Path
from typing import Callable, Optional

from .models import (
    ConversionError,
    ConversionOptions,
    ConversionProgress,
    ConversionRequest,
    ConversionResult,
    ConversionStatus,
    OutputFormat,
)


class ToMarkdownConverter:
    """Converts supported document formats to Markdown using MarkItDown."""

    def __init__(self) -> None:
        self._md: Optional[object] = None

    def _get_engine(self):
        """Lazy-load MarkItDown to avoid startup cost."""
        if self._md is None:
            from markitdown import MarkItDown
            self._md = MarkItDown()
        return self._md

    def convert(
        self,
        request: ConversionRequest,
        progress_callback: Optional[Callable[[ConversionProgress], None]] = None,
    ) -> ConversionResult | ConversionError:
        import time

        start = time.monotonic()
        filename = request.input_path.name

        def emit(percent: int, status: ConversionStatus, message: str = "") -> None:
            if progress_callback:
                progress_callback(
                    ConversionProgress(
                        request_id=request.request_id,
                        filename=filename,
                        percent=percent,
                        status=status,
                        message=message,
                    )
                )

        emit(0, ConversionStatus.CONVERTING, "Starting…")

        try:
            # Validate input
            if not request.input_path.exists():
                raise FileNotFoundError(f"File not found: {request.input_path}")

            emit(10, ConversionStatus.CONVERTING, "Reading file…")

            markdown_content: str = ""
            engine_error: Optional[Exception] = None

            try:
                engine = self._get_engine()
                emit(30, ConversionStatus.CONVERTING, "Converting…")
                result = engine.convert(str(request.input_path))
                markdown_content = result.text_content or ""
            except Exception as exc:
                engine_error = exc

            # If primary engine returned empty or failed, invoke resilient format-specific fallback
            if not markdown_content.strip():
                fallback_text = self._fallback_convert(request.input_path)
                if fallback_text:
                    markdown_content = fallback_text
                elif engine_error:
                    raise engine_error

            emit(80, ConversionStatus.CONVERTING, "Writing output…")

            # Ensure output directory exists
            output_path = request.output_path
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(markdown_content, encoding="utf-8")

            duration_ms = int((time.monotonic() - start) * 1000)
            emit(100, ConversionStatus.DONE, "Done")

            return ConversionResult(
                request=request,
                output_path=output_path,
                success=True,
                duration_ms=duration_ms,
            )

        except Exception as exc:
            emit(0, ConversionStatus.ERROR, "Failed")
            return ConversionError(
                request=request,
                user_message=self._friendly_message(exc, request.input_path),
                detail=traceback.format_exc(),
            )

    def _fallback_convert(self, path: Path) -> Optional[str]:
        """Resilient fallback converter when the primary engine is unavailable or fails."""
        ext = path.suffix.lower()
        try:
            if ext == ".docx":
                return self._fallback_docx(path)
            elif ext == ".pdf":
                return self._fallback_pdf(path)
            elif ext in (".txt", ".text", ".log"):
                return self._fallback_txt(path)
            elif ext == ".csv":
                return self._fallback_csv(path)
            elif ext in (".html", ".htm"):
                return self._fallback_html(path)
            elif ext in (".xlsx", ".xls"):
                return self._fallback_xlsx(path)
        except Exception:
            return None
        return None

    @staticmethod
    def _fallback_docx(path: Path) -> str:
        import docx

        doc = docx.Document(str(path))
        lines: list[str] = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            style_name = (p.style.name or "").lower() if p.style else ""
            if "heading 1" in style_name:
                lines.append(f"# {text}\n")
            elif "heading 2" in style_name:
                lines.append(f"## {text}\n")
            elif "heading 3" in style_name:
                lines.append(f"### {text}\n")
            elif "list" in style_name or "bullet" in style_name:
                lines.append(f"- {text}")
            else:
                lines.append(f"{text}\n")

        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                table_rows.append([c.text.strip().replace("\n", " ") for c in row.cells])
            if table_rows:
                col_count = len(table_rows[0])
                header = "| " + " | ".join(table_rows[0]) + " |"
                sep = "| " + " | ".join(["---"] * col_count) + " |"
                lines.append(f"\n{header}\n{sep}")
                for r in table_rows[1:]:
                    padded = r + [""] * (col_count - len(r))
                    lines.append("| " + " | ".join(padded[:col_count]) + " |")
                lines.append("")

        return "\n".join(lines).strip()

    @staticmethod
    def _fallback_pdf(path: Path) -> str:
        # Try pypdf first (pure python, zero external C-dependencies)
        try:
            import pypdf

            reader = pypdf.PdfReader(str(path))
            chunks = []
            for page in reader.pages:
                text = page.extract_text()
                if text and text.strip():
                    chunks.append(text.strip())
            if chunks:
                return "\n\n".join(chunks)
        except Exception:
            pass

        # Try pdfminer
        try:
            import pdfminer.high_level

            text = pdfminer.high_level.extract_text(str(path))
            if text and text.strip():
                return text.strip()
        except Exception:
            pass

        # Try pdfplumber
        try:
            import pdfplumber

            chunks = []
            with pdfplumber.open(str(path)) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text and text.strip():
                        chunks.append(text.strip())
            if chunks:
                return "\n\n".join(chunks)
        except Exception:
            pass

        return ""

    @staticmethod
    def _fallback_txt(path: Path) -> str:
        for enc in ("utf-8", "utf-8-sig", "latin1", "cp1252"):
            try:
                return path.read_text(encoding=enc)
            except UnicodeDecodeError:
                continue
        return path.read_text(encoding="utf-8", errors="replace")

    @staticmethod
    def _fallback_csv(path: Path) -> str:
        import csv

        rows = []
        for enc in ("utf-8", "utf-8-sig", "latin1", "cp1252"):
            try:
                with open(path, "r", encoding=enc, newline="") as f:
                    rows = list(csv.reader(f))
                break
            except UnicodeDecodeError:
                continue
        if not rows:
            with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
                rows = list(csv.reader(f))
        if not rows:
            return ""

        col_count = max(len(r) for r in rows)
        lines = [
            "| " + " | ".join(rows[0] + [""] * (col_count - len(rows[0]))) + " |",
            "| " + " | ".join(["---"] * col_count) + " |",
        ]
        for r in rows[1:]:
            padded = r + [""] * (col_count - len(r))
            lines.append("| " + " | ".join(padded[:col_count]) + " |")
        return "\n".join(lines)

    @staticmethod
    def _fallback_html(path: Path) -> str:
        html = path.read_text(encoding="utf-8", errors="replace")
        try:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(html, "html.parser")
            return soup.get_text("\n\n").strip()
        except Exception:
            import re

            return re.sub(r"<[^>]+>", "", html).strip()

    @staticmethod
    def _fallback_xlsx(path: Path) -> str:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(str(path), data_only=True)
            sheet = wb.active
            lines = []
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    lines.append("| " + " | ".join(str(c or "") for c in row) + " |")
            return "\n".join(lines)
        except Exception:
            return ""

    @staticmethod
    def _friendly_message(exc: Exception, path: Path) -> str:
        name = path.name
        ext = path.suffix.upper().lstrip(".")

        if isinstance(exc, FileNotFoundError):
            return f'The file "{name}" could not be found.'
        if isinstance(exc, PermissionError):
            return f'Permission denied when reading "{name}". Is it open in another application?'
        if "password" in str(exc).lower() or "encrypted" in str(exc).lower():
            return f'"{name}" appears to be password-protected.'
        if "corrupted" in str(exc).lower() or "invalid" in str(exc).lower():
            return f'"{name}" appears to be damaged or is not a valid {ext} file.'

        return (
            f'Couldn\'t convert "{name}".\n'
            f"The {ext} file could not be processed. It may be damaged, "
            f"unsupported, or password-protected."
        )

