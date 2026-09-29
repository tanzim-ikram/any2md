"""Fidelity regression tests for DOCX -> Markdown and PDF -> Markdown.

These build their fixtures at runtime (python-docx / reportlab, both already
project dependencies) so no binary files need to be committed.
"""

from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

import pytest

from any2md.conversion.engine import ConversionEngine
from any2md.conversion.models import (
    ConversionOptions,
    ConversionRequest,
    ConversionResult,
    OutputFormat,
)


def _tiny_png_bytes() -> bytes:
    raw = b"\x00\xff\x00\x00"

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def _make_docx(path: Path, with_image: bool) -> None:
    import docx

    doc = docx.Document()
    doc.add_paragraph("Doc Title", style="Title")
    doc.add_heading("Sub Heading", level=2)
    doc.add_paragraph("A wise quote", style="Quote")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Sales"
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "100"
    if with_image:
        doc.add_picture(io.BytesIO(_tiny_png_bytes()))
    doc.save(str(path))


def _make_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Table, TableStyle

    styles = getSampleStyleSheet()
    table = Table([["Region", "Sales", "Growth"], ["North", "100", "5%"], ["South", "200", "7%"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))

    SimpleDocTemplate(str(path)).build(
        [
            Paragraph("Annual Report", styles["Title"]),
            Paragraph("Intro with <b>bold</b> and <i>italic</i> text.", styles["Normal"]),
            Paragraph("Results", styles["Heading2"]),
            ListFlowable(
                [Paragraph("First point", styles["Normal"]), Paragraph("Second point", styles["Normal"])],
                bulletType="bullet",
            ),
            table,
            Paragraph("Conclusion", styles["Heading2"]),
            Paragraph("Done.", styles["Normal"]),
        ]
    )


@pytest.fixture
def engine() -> ConversionEngine:
    return ConversionEngine()


class TestDocxFidelity:
    def test_title_and_quote_styles_become_headings_and_blockquotes(self, tmp_path: Path, engine: ConversionEngine) -> None:
        docx_path = tmp_path / "sample.docx"
        _make_docx(docx_path, with_image=False)

        req = ConversionRequest(input_path=docx_path, output_format=OutputFormat.MARKDOWN, output_dir=tmp_path)
        res = engine.convert(req)
        assert isinstance(res, ConversionResult)

        content = res.output_path.read_text(encoding="utf-8")
        assert "# Doc Title" in content
        assert "> A wise quote" in content

    def test_table_header_row_is_not_empty(self, tmp_path: Path, engine: ConversionEngine) -> None:
        docx_path = tmp_path / "sample.docx"
        _make_docx(docx_path, with_image=False)

        req = ConversionRequest(input_path=docx_path, output_format=OutputFormat.MARKDOWN, output_dir=tmp_path)
        res = engine.convert(req)
        assert isinstance(res, ConversionResult)

        content = res.output_path.read_text(encoding="utf-8")
        assert "| Region | Sales |" in content
        assert "|  |  |" not in content

    def test_image_kept_when_preserve_images_enabled(self, tmp_path: Path, engine: ConversionEngine) -> None:
        docx_path = tmp_path / "sample.docx"
        _make_docx(docx_path, with_image=True)

        req = ConversionRequest(
            input_path=docx_path,
            output_format=OutputFormat.MARKDOWN,
            output_dir=tmp_path,
            options=ConversionOptions(preserve_images=True),
        )
        res = engine.convert(req)
        assert isinstance(res, ConversionResult)
        assert "data:image/png;base64," in res.output_path.read_text(encoding="utf-8")

    def test_image_dropped_when_preserve_images_disabled(self, tmp_path: Path, engine: ConversionEngine) -> None:
        docx_path = tmp_path / "sample.docx"
        _make_docx(docx_path, with_image=True)

        req = ConversionRequest(
            input_path=docx_path,
            output_format=OutputFormat.MARKDOWN,
            output_dir=tmp_path,
            options=ConversionOptions(preserve_images=False),
        )
        res = engine.convert(req)
        assert isinstance(res, ConversionResult)
        assert "base64" not in res.output_path.read_text(encoding="utf-8")


class TestPdfFidelity:
    def test_headings_and_table_and_emphasis(self, tmp_path: Path, engine: ConversionEngine) -> None:
        pdf_path = tmp_path / "report.pdf"
        _make_pdf(pdf_path)

        req = ConversionRequest(input_path=pdf_path, output_format=OutputFormat.MARKDOWN, output_dir=tmp_path)
        res = engine.convert(req)
        assert isinstance(res, ConversionResult)

        content = res.output_path.read_text(encoding="utf-8")
        assert "# Annual Report" in content
        assert "## Results" in content
        assert "**bold**" in content
        assert "- First point" in content
        assert "- Second point" in content
        assert "(cid:" not in content

        # Table rows/columns must stay together, not be scattered by column.
        assert "| Region | Sales | Growth |" in content
        north_line = next(line for line in content.splitlines() if "North" in line)
        assert "100" in north_line and "5%" in north_line


class TestTwoHopImageSafety:
    def test_pdf_to_docx_does_not_embed_base64_text(self, tmp_path: Path, engine: ConversionEngine) -> None:
        import docx

        pdf_path = tmp_path / "report.pdf"
        _make_pdf(pdf_path)

        req = ConversionRequest(
            input_path=pdf_path,
            output_format=OutputFormat.DOCX,
            output_dir=tmp_path,
            options=ConversionOptions(preserve_images=True),
        )
        res = engine.convert(req)
        assert isinstance(res, ConversionResult)

        doc = docx.Document(str(res.output_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "base64" not in full_text
